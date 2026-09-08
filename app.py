# -*- coding: utf-8 -*-
from flask import Flask, render_template, request, jsonify, send_file, session, redirect, url_for
import sqlite3
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta
from werkzeug.security import generate_password_hash, check_password_hash
import io
import os
import secrets
import functools
import json
import threading

# Load .env if present (local dev); on hosting platforms env vars are set natively
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

app = Flask(__name__)

_data_version = 0
_dv_lock = threading.Lock()

# Simple in-memory login rate limiter: {ip: [timestamp, ...]}
_login_attempts: dict = {}
_login_lock = threading.Lock()
_MAX_ATTEMPTS = 10
_WINDOW_SECS  = 900  # 15 minutes

def _check_rate_limit(ip: str) -> bool:
    """Return True if the IP is allowed to attempt login."""
    now = datetime.now().timestamp()
    with _login_lock:
        attempts = [t for t in _login_attempts.get(ip, []) if now - t < _WINDOW_SECS]
        _login_attempts[ip] = attempts
        return len(attempts) < _MAX_ATTEMPTS

def _record_failed_login(ip: str):
    now = datetime.now().timestamp()
    with _login_lock:
        _login_attempts.setdefault(ip, []).append(now)

def _clear_login_attempts(ip: str):
    with _login_lock:
        _login_attempts.pop(ip, None)

def bump_version():
    global _data_version
    with _dv_lock:
        _data_version += 1

_base = os.path.dirname(__file__)
DB_PATH = os.path.join(_base, os.environ.get('DATABASE_PATH', 'mis.db'))

# Secret key: prefer env var, fall back to persisted file for local dev
_secret = os.environ.get('SECRET_KEY', '')
if not _secret:
    _sk_path = os.path.join(_base, '.secret_key')
    if os.path.exists(_sk_path):
        with open(_sk_path) as _f:
            _secret = _f.read().strip()
    if not _secret:
        _secret = secrets.token_hex(32)
        with open(_sk_path, 'w') as _f:
            _f.write(_secret)
app.secret_key = _secret
app.permanent_session_lifetime = timedelta(hours=12)

# Ordered list of every editable field (no id, sr_no, sort_order, created_at)
FIELDS = [
    'cat', 'company_name', 'activity', 'country',
    'dnb_rep_validity', 'dnb_rating', 'dnb_processed_by',
    'audit_rep', 'audit_processed_by', 'audit_prepared_by',
    'mgtm_ac_q1', 'mgtm_ac_q2', 'mgtm_ac_q3', 'mgtm_ac_q4',
    'aecb_dir', 'aecb_com', 'cibil_dir',
    'remarks',
]


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


# ── Auth helpers ──────────────────────────────────────────────────────────────
def require_auth(f):
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            if request.path.startswith('/api/') or request.path.startswith('/export'):
                return jsonify({'error': 'Unauthorized'}), 401
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated


def init_db():
    with get_db() as conn:
        conn.execute('''
            CREATE TABLE IF NOT EXISTS users (
                id            INTEGER PRIMARY KEY AUTOINCREMENT,
                name          TEXT NOT NULL,
                email         TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                role          TEXT DEFAULT 'user',
                approved      INTEGER DEFAULT 0,
                created_at    TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS logs (
                id         INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id    INTEGER,
                user_name  TEXT,
                user_role  TEXT,
                action     TEXT NOT NULL,
                details    TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS settings (
                key   TEXT PRIMARY KEY,
                value TEXT NOT NULL
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS password_resets (
                id         INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id    INTEGER NOT NULL,
                token      TEXT UNIQUE NOT NULL,
                expires_at TEXT NOT NULL,
                used       INTEGER DEFAULT 0
            )
        ''')
        conn.execute('''
            CREATE TABLE IF NOT EXISTS companies (
                id            INTEGER PRIMARY KEY AUTOINCREMENT,
                cat           TEXT,
                company_name  TEXT NOT NULL,
                activity      TEXT,
                country       TEXT DEFAULT 'UAE',
                dnb_rep_validity   TEXT,
                dnb_rating         TEXT,
                dnb_processed_by   TEXT,
                audit_rep          TEXT,
                audit_processed_by TEXT,
                audit_prepared_by  TEXT,
                mgtm_ac_q1    TEXT,
                mgtm_ac_q2    TEXT,
                mgtm_ac_q3    TEXT,
                mgtm_ac_q4    TEXT,
                aecb_dir      TEXT,
                aecb_com      TEXT,
                cibil_dir     TEXT,
                remarks       TEXT,
                sort_order    INTEGER DEFAULT 0,
                created_at    TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        conn.commit()
    _migrate()


def _migrate():
    """Add any columns that don't exist yet (handles existing DBs)."""
    new_cols = [
        ('dnb_processed_by',   'TEXT'),
        ('audit_processed_by', 'TEXT'),
        ('audit_prepared_by',  'TEXT'),
        ('mgtm_ac_q1', 'TEXT'),
        ('mgtm_ac_q2', 'TEXT'),
        ('mgtm_ac_q3', 'TEXT'),
        ('mgtm_ac_q4', 'TEXT'),
    ]
    user_cols = [
        ('role',        "TEXT DEFAULT 'user'"),
        ('approved',    'INTEGER DEFAULT 0'),
        ('linked_name', 'TEXT'),
    ]
    with get_db() as conn:
        existing = {row[1] for row in conn.execute('PRAGMA table_info(companies)').fetchall()}
        for col_name, col_type in new_cols:
            if col_name not in existing:
                conn.execute(f'ALTER TABLE companies ADD COLUMN {col_name} {col_type}')
        existing_u = {row[1] for row in conn.execute('PRAGMA table_info(users)').fetchall()}
        for col_name, col_type in user_cols:
            if col_name not in existing_u:
                conn.execute(f'ALTER TABLE users ADD COLUMN {col_name} {col_type}')
        conn.commit()


# ── Auth routes ───────────────────────────────────────────────────────────────
@app.route('/login', methods=['GET', 'POST'])
def login():
    if 'user_id' in session:
        return redirect('/')
    if request.method == 'POST':
        ip = request.remote_addr
        if not _check_rate_limit(ip):
            return jsonify({'error': 'Too many failed attempts. Try again in 15 minutes.'}), 429
        data = request.json or {}
        email = data.get('email', '').strip().lower()
        password = data.get('password', '')
        with get_db() as conn:
            row = conn.execute('SELECT id, name, password_hash, role, approved FROM users WHERE email=?', (email,)).fetchone()
        if row and check_password_hash(row['password_hash'], password):
            if not row['approved']:
                return jsonify({'error': 'Your account is pending approval by an admin.'}), 403
            _clear_login_attempts(ip)
            session.permanent = True
            session['user_id'] = row['id']
            session['user_name'] = row['name']
            session['user_role'] = row['role']
            log_action('Login', f'User "{row["name"]}" ({email}) signed in as {row["role"]} from IP {request.remote_addr}')
            return jsonify({'ok': True})
        _record_failed_login(ip)
        return jsonify({'error': 'Invalid email or password'}), 401
    return render_template('login.html')


@app.route('/register', methods=['GET', 'POST'])
def register():
    if 'user_id' in session:
        return redirect('/')
    if request.method == 'POST':
        data = request.json or {}
        name = data.get('name', '').strip()
        email = data.get('email', '').strip().lower()
        password = data.get('password', '')
        if not name or not email or len(password) < 6:
            return jsonify({'error': 'All fields required. Password min 6 characters.'}), 400
        try:
            with get_db() as conn:
                count = conn.execute('SELECT COUNT(*) FROM users').fetchone()[0]
                role = 'superadmin' if count == 0 else 'user'
                approved = 1 if count == 0 else 0
                conn.execute(
                    'INSERT INTO users (name, email, password_hash, role, approved) VALUES (?,?,?,?,?)',
                    (name, email, generate_password_hash(password), role, approved)
                )
                conn.commit()
            if role == 'superadmin':
                return jsonify({'ok': True})
            return jsonify({'ok': True, 'pending': True})
        except sqlite3.IntegrityError:
            return jsonify({'error': 'An account with this email already exists.'}), 400
    return render_template('register.html')


@app.route('/logout')
def logout():
    log_action('Logout', f'User "{session.get("user_name")}" ({session.get("user_role")}) signed out')
    session.clear()
    return redirect('/login')


@app.route('/forgot-password', methods=['GET', 'POST'])
def forgot_password():
    if request.method == 'POST':
        if not _check_rate_limit(request.remote_addr):
            return jsonify({'ok': True})  # silent throttle, don't reveal rate limit
        data = request.json or {}
        email = data.get('email', '').strip().lower()
        with get_db() as conn:
            row = conn.execute('SELECT id, name FROM users WHERE email=?', (email,)).fetchone()
            smtp_rows = conn.execute(
                "SELECT key, value FROM settings WHERE key IN ('smtp_email','smtp_password')"
            ).fetchall()
        smtp = {r['key']: r['value'] for r in smtp_rows}
        smtp_email = smtp.get('smtp_email', '').strip()
        smtp_pass  = smtp.get('smtp_password', '').strip()

        if not row:
            # Always return ok to avoid user enumeration
            return jsonify({'ok': True, 'sent': bool(smtp_email)})

        token = secrets.token_urlsafe(32)
        expires = (datetime.now() + timedelta(hours=2)).isoformat()
        with get_db() as conn:
            conn.execute('DELETE FROM password_resets WHERE user_id=?', (row['id'],))
            conn.execute('INSERT INTO password_resets (user_id, token, expires_at) VALUES (?,?,?)',
                         (row['id'], token, expires))
            conn.commit()
        reset_url = url_for('reset_password', token=token, _external=True)

        if smtp_email and smtp_pass:
            try:
                msg = MIMEMultipart('alternative')
                msg['Subject'] = 'MIS — Reset your password'
                msg['From'] = smtp_email
                msg['To'] = email
                body = (
                    f'Hi {row["name"]},\n\n'
                    f'You requested a password reset for the MIS portal.\n\n'
                    f'Click the link below to set a new password (valid for 2 hours):\n{reset_url}\n\n'
                    f'If you did not request this, you can safely ignore this email.\n\n'
                    f'— MIS System'
                )
                msg.attach(MIMEText(body, 'plain'))
                with smtplib.SMTP_SSL('smtp.gmail.com', 465) as server:
                    server.login(smtp_email, smtp_pass)
                    server.sendmail(smtp_email, email, msg.as_string())
            except Exception:
                return jsonify({'error': 'Failed to send email. Check your SMTP settings.'}), 500
            return jsonify({'ok': True, 'sent': True})
        else:
            # SMTP not configured — return URL for admin to share manually
            return jsonify({'ok': True, 'sent': False, 'reset_url': reset_url})
    return render_template('forgot_password.html')


@app.route('/reset-password/<token>', methods=['GET', 'POST'])
def reset_password(token):
    if request.method == 'POST':
        data = request.json or {}
        password = data.get('password', '')
        if len(password) < 6:
            return jsonify({'error': 'Password must be at least 6 characters.'}), 400
        with get_db() as conn:
            row = conn.execute(
                'SELECT user_id FROM password_resets WHERE token=? AND used=0 AND expires_at>?',
                (token, datetime.now().isoformat())
            ).fetchone()
            if not row:
                return jsonify({'error': 'Reset link is invalid or has expired.'}), 400
            conn.execute('UPDATE users SET password_hash=? WHERE id=?',
                         (generate_password_hash(password), row['user_id']))
            conn.execute('UPDATE password_resets SET used=1 WHERE token=?', (token,))
            conn.commit()
        return jsonify({'ok': True})
    return render_template('reset_password.html', token=token)


@app.route('/api/me')
def me():
    if 'user_id' not in session:
        return jsonify({'error': 'Unauthorized'}), 401
    with get_db() as conn:
        row = conn.execute('SELECT name, role, linked_name FROM users WHERE id=?', (session['user_id'],)).fetchone()
    if not row:
        return jsonify({'error': 'Unauthorized'}), 401
    return jsonify({'id': session['user_id'], 'name': row['name'], 'role': row['role'], 'linked_name': row['linked_name'] or ''})


def _get_live_role():
    """Fetch the user's current role from DB and sync it into session."""
    uid = session.get('user_id')
    if not uid:
        return None
    with get_db() as conn:
        row = conn.execute('SELECT role FROM users WHERE id=?', (uid,)).fetchone()
    if not row:
        session.clear()
        return None
    role = row['role']
    session['user_role'] = role  # keep session in sync
    return role


def require_admin(f):
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return jsonify({'error': 'Unauthorized'}), 401
        role = _get_live_role()
        if role not in ('admin', 'superadmin'):
            return jsonify({'error': 'Forbidden'}), 403
        return f(*args, **kwargs)
    return decorated


def require_superadmin(f):
    @functools.wraps(f)
    def decorated(*args, **kwargs):
        if 'user_id' not in session:
            return jsonify({'error': 'Unauthorized'}), 401
        role = _get_live_role()
        if role != 'superadmin':
            return jsonify({'error': 'Forbidden'}), 403
        return f(*args, **kwargs)
    return decorated


@app.route('/api/admin/users')
@require_admin
def admin_list_users():
    with get_db() as conn:
        rows = conn.execute(
            'SELECT id, name, email, role, approved, linked_name, created_at FROM users ORDER BY created_at'
        ).fetchall()
    return jsonify([dict(r) for r in rows])


@app.route('/api/admin/users/<int:uid>/approve', methods=['POST'])
@require_admin
def admin_approve_user(uid):
    with get_db() as conn:
        row = conn.execute('SELECT name, email FROM users WHERE id=?', (uid,)).fetchone()
        conn.execute('UPDATE users SET approved=1 WHERE id=?', (uid,))
        conn.commit()
    bump_version()
    log_action('User Approved', f'Approved account for "{row["name"]}" ({row["email"]}). Account is now active and user can log in.' if row else f'Approved user ID {uid}')
    return jsonify({'ok': True})


@app.route('/api/admin/users/<int:uid>', methods=['PATCH'])
@require_admin
def admin_update_user(uid):
    is_superadmin = session.get('user_role') == 'superadmin'
    if not is_superadmin and uid != session.get('user_id'):
        return jsonify({'error': 'Admins can only edit their own profile'}), 403
    data = request.json or {}
    fields, values = [], []
    if 'name' in data:
        name = data['name'].strip()
        if not name:
            return jsonify({'error': 'Name cannot be empty'}), 400
        fields.append('name=?'); values.append(name)
    if 'linked_name' in data:
        fields.append('linked_name=?'); values.append(data['linked_name'] or None)
    if not fields:
        return jsonify({'error': 'Nothing to update'}), 400
    values.append(uid)
    with get_db() as conn:
        target = conn.execute('SELECT name, email FROM users WHERE id=?', (uid,)).fetchone()
        conn.execute(f'UPDATE users SET {", ".join(fields)} WHERE id=?', values)
        conn.commit()
    tname = target['name'] if target else f'ID {uid}'
    temail = target['email'] if target else ''
    if 'name' in data:
        log_action('User Renamed', f'Renamed account "{tname}" ({temail}) to "{data["name"].strip()}"')
    if 'linked_name' in data:
        linked = data['linked_name'] or '(none)'
        log_action('Staff Name Linked', f'Linked account "{tname}" ({temail}) to staff name "{linked}". Task notifications and alerts will now use this name.')
    return jsonify({'ok': True})


@app.route('/api/admin/users/<int:uid>/reject', methods=['POST'])
@require_superadmin
def admin_reject_user(uid):
    if uid == session.get('user_id'):
        return jsonify({'error': 'Cannot remove yourself'}), 400
    with get_db() as conn:
        row = conn.execute('SELECT name, email FROM users WHERE id=?', (uid,)).fetchone()
        conn.execute('DELETE FROM users WHERE id=?', (uid,))
        conn.commit()
    bump_version()
    if row:
        log_action('User Removed', f'Permanently deleted account for "{row["name"]}" ({row["email"]}). All access has been revoked.')
    else:
        log_action('User Removed', f'Deleted user ID {uid}')
    return jsonify({'ok': True})


@app.route('/api/settings', methods=['GET'])
@require_auth
def get_settings():
    is_superadmin = _get_live_role() == 'superadmin'
    with get_db() as conn:
        rows = conn.execute('SELECT key, value FROM settings').fetchall()
    SENSITIVE = {'smtp_password', 'smtp_email'}
    result = {}
    for r in rows:
        if r['key'] in SENSITIVE and not is_superadmin:
            continue
        if r['key'] == 'smtp_password':
            continue  # never send password to client; existence is enough
        result[r['key']] = r['value']
    return jsonify(result)


SMTP_KEYS = {'smtp_email', 'smtp_password'}

@app.route('/api/settings', methods=['PUT'])
@require_auth
def save_settings():
    data = request.json or {}
    live_role = _get_live_role()
    is_superadmin = live_role == 'superadmin'
    is_admin = live_role in ('admin', 'superadmin')
    # SMTP keys require superadmin
    if any(k in SMTP_KEYS for k in data) and not is_superadmin:
        return jsonify({'error': 'Forbidden'}), 403
    # All other settings require at least admin
    if not is_admin:
        # Regular users may only write list settings (staff, countries, etc.)
        LIST_KEYS = {'staff', 'countries', 'activities', 'ratings', 'aecb'}
        if any(k not in LIST_KEYS for k in data):
            return jsonify({'error': 'Forbidden'}), 403
    with get_db() as conn:
        for key, value in data.items():
            conn.execute('INSERT INTO settings (key, value) VALUES (?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value',
                         (key, value))
        conn.commit()
    return jsonify({'ok': True})


# ── Logging ───────────────────────────────────────────────────────────────────
def log_action(action, details=''):
    if 'user_id' not in session:
        return
    try:
        with get_db() as conn:
            row = conn.execute('SELECT name, role FROM users WHERE id=?', (session['user_id'],)).fetchone()
            user_name = row['name'] if row else session.get('user_name', '')
            user_role = row['role'] if row else session.get('user_role', '')
            conn.execute(
                'INSERT INTO logs (user_id, user_name, user_role, action, details) VALUES (?,?,?,?,?)',
                (session['user_id'], user_name, user_role, action, details)
            )
            conn.commit()
    except Exception:
        pass


@app.route('/api/logs')
@require_admin
def get_logs():
    with get_db() as conn:
        rows = conn.execute(
            'SELECT id, user_name, user_role, action, details, created_at FROM logs ORDER BY created_at DESC LIMIT 500'
        ).fetchall()
    return jsonify([dict(r) for r in rows])


@app.route('/api/logs', methods=['DELETE'])
@require_superadmin
def clear_logs():
    with get_db() as conn:
        count = conn.execute('SELECT COUNT(*) FROM logs').fetchone()[0]
        conn.execute('DELETE FROM logs')
        conn.commit()
    log_action('Logs Cleared', f'Cleared all {count} log entries from the audit trail. This action cannot be undone.')
    return jsonify({'ok': True})


@app.route('/api/logs/<int:log_id>', methods=['DELETE'])
@require_superadmin
def delete_log(log_id):
    with get_db() as conn:
        conn.execute('DELETE FROM logs WHERE id=?', (log_id,))
        conn.commit()
    return jsonify({'ok': True})


@app.route('/api/admin/users/<int:uid>/role', methods=['POST'])
@require_superadmin
def admin_set_role(uid):
    if uid == session.get('user_id'):
        return jsonify({'error': 'Cannot change your own role'}), 400
    data = request.json or {}
    role = data.get('role', '')
    if role not in ('user', 'admin', 'superadmin'):
        return jsonify({'error': 'Invalid role'}), 400
    with get_db() as conn:
        row = conn.execute('SELECT name, email, role FROM users WHERE id=?', (uid,)).fetchone()
        conn.execute('UPDATE users SET role=? WHERE id=?', (role, uid))
        conn.commit()
    if row:
        log_action('Role Changed', f'Changed role for "{row["name"]}" ({row["email"]}) from "{row["role"]}" to "{role}"')
    else:
        log_action('Role Changed', f'Changed role for user ID {uid} to "{role}"')
    return jsonify({'ok': True})


@app.route('/api/version')
@require_auth
def data_version():
    return jsonify({'v': _data_version})


# ── Backup / Restore ──────────────────────────────────────────────────────────
@app.route('/api/backup')
@require_superadmin
def backup():
    with get_db() as conn:
        companies = [dict(r) for r in conn.execute('SELECT * FROM companies ORDER BY id').fetchall()]
        users     = [dict(r) for r in conn.execute('SELECT * FROM users ORDER BY id').fetchall()]
        settings  = [dict(r) for r in conn.execute('SELECT * FROM settings').fetchall()]
        logs      = [dict(r) for r in conn.execute('SELECT * FROM logs ORDER BY id').fetchall()]
    payload = {
        'version': 1,
        'exported_at': datetime.now().isoformat(),
        'tables': {'companies': companies, 'users': users, 'settings': settings, 'logs': logs},
    }
    log_action('Backup Created', f'Full database backup downloaded ({len(companies)} companies, {len(users)} users, {len(logs)} log entries)')
    buf = io.BytesIO(json.dumps(payload, indent=2, default=str).encode('utf-8'))
    buf.seek(0)
    filename = f'MIS_backup_{datetime.now().strftime("%Y%m%d_%H%M%S")}.json'
    return send_file(buf, as_attachment=True, download_name=filename, mimetype='application/json')


@app.route('/api/restore', methods=['POST'])
@require_superadmin
def restore():
    f = request.files.get('backup')
    if not f:
        return jsonify({'error': 'No file uploaded'}), 400
    try:
        data = json.load(f)
    except Exception:
        return jsonify({'error': 'Invalid backup file — must be a valid JSON backup'}), 400
    if 'tables' not in data or 'version' not in data:
        return jsonify({'error': 'Unrecognised backup format'}), 400

    tables = data['tables']
    ALLOWED = {
        'companies': ['id','cat','company_name','activity','country','dnb_rep_validity','dnb_rating',
                      'dnb_processed_by','audit_rep','audit_processed_by','audit_prepared_by',
                      'mgtm_ac_q1','mgtm_ac_q2','mgtm_ac_q3','mgtm_ac_q4','aecb_dir','aecb_com',
                      'cibil_dir','remarks','sort_order','created_at'],
        'users':     ['id','name','email','password_hash','role','approved','linked_name','created_at'],
        'settings':  ['key','value'],
        'logs':      ['id','user_id','user_name','user_role','action','details','created_at'],
    }

    def safe_insert(conn, table, rows):
        allowed = ALLOWED[table]
        conn.execute(f'DELETE FROM {table}')
        for row in rows:
            clean = {k: v for k, v in row.items() if k in allowed}
            if not clean:
                continue
            cols = ', '.join(clean.keys())
            placeholders = ', '.join(['?'] * len(clean))
            conn.execute(f'INSERT INTO {table} ({cols}) VALUES ({placeholders})', list(clean.values()))

    with get_db() as conn:
        for tbl in ('companies', 'users', 'settings', 'logs'):
            if tbl in tables:
                safe_insert(conn, tbl, tables[tbl])
        conn.commit()

    bump_version()
    exported_at = data.get('exported_at', 'unknown')
    log_action('Backup Restored', f'Database fully restored from backup exported at {exported_at}')
    return jsonify({'ok': True})


@app.route('/')
@require_auth
def index():
    return render_template('index.html')


@app.route('/api/companies', methods=['GET'])
@require_auth
def get_companies():
    with get_db() as conn:
        rows = conn.execute(
            'SELECT * FROM companies ORDER BY country, cat, sort_order, id'
        ).fetchall()

    companies = []
    sr_by_country: dict = {}
    for row in rows:
        d = dict(row)
        country = (d.get('country') or 'UAE').strip().upper()
        sr_by_country.setdefault(country, 1)
        d['sr_no'] = sr_by_country[country]
        sr_by_country[country] += 1
        companies.append(d)

    return jsonify(companies)


@app.route('/api/companies', methods=['POST'])
@require_admin
def add_company():
    data = request.json
    placeholders = ', '.join(['?'] * len(FIELDS))
    cols = ', '.join(FIELDS)
    values = [data.get(f) or None for f in FIELDS]

    with get_db() as conn:
        cur = conn.execute(
            f'INSERT INTO companies ({cols}, sort_order) VALUES ({placeholders}, ?)',
            values + [data.get('sort_order', 0)],
        )
        conn.commit()
        row = conn.execute('SELECT * FROM companies WHERE id=?', (cur.lastrowid,)).fetchone()
    bump_version()
    FIELD_LABELS = {
        'cat':'Category','company_name':'Company Name','activity':'Activity','country':'Country',
        'dnb_rep_validity':'D&B Rep Validity','dnb_rating':'D&B Rating','dnb_processed_by':'D&B Processed By',
        'audit_rep':'Last Audit','audit_processed_by':'Audit Processed By','audit_prepared_by':'Auditor Name',
        'mgtm_ac_q1':'Mgtm A/c Q1','mgtm_ac_q2':'Mgtm A/c Q2','mgtm_ac_q3':'Mgtm A/c Q3','mgtm_ac_q4':'Mgtm A/c Q4',
        'aecb_dir':'AECB (Dir)','aecb_com':'AECB (Com)','cibil_dir':'Cibil (Dir)','remarks':'Remarks',
    }
    filled = [f'{FIELD_LABELS.get(f,f)}: "{data.get(f)}"' for f in FIELDS if data.get(f)]
    details = f'Added "{data.get("company_name")}" ({data.get("country","UAE")}). Fields: {" | ".join(filled)}'
    log_action('Company Added', details)
    return jsonify(dict(row)), 201


@app.route('/api/companies/<int:company_id>', methods=['PUT'])
@require_admin
def update_company(company_id):
    data = request.json
    set_clause = ', '.join(f'{f}=?' for f in FIELDS)
    values = [data.get(f) or None for f in FIELDS]

    FIELD_LABELS = {
        'cat':'Category','company_name':'Company Name','activity':'Activity','country':'Country',
        'dnb_rep_validity':'D&B Rep Validity','dnb_rating':'D&B Rating','dnb_processed_by':'D&B Processed By',
        'audit_rep':'Last Audit','audit_processed_by':'Audit Processed By','audit_prepared_by':'Auditor Name',
        'mgtm_ac_q1':'Mgtm A/c Q1','mgtm_ac_q2':'Mgtm A/c Q2','mgtm_ac_q3':'Mgtm A/c Q3','mgtm_ac_q4':'Mgtm A/c Q4',
        'aecb_dir':'AECB (Dir)','aecb_com':'AECB (Com)','cibil_dir':'Cibil (Dir)','remarks':'Remarks',
    }

    with get_db() as conn:
        old = dict(conn.execute('SELECT * FROM companies WHERE id=?', (company_id,)).fetchone() or {})
        conn.execute(
            f'UPDATE companies SET {set_clause}, sort_order=? WHERE id=?',
            values + [data.get('sort_order', 0), company_id],
        )
        conn.commit()
        row = conn.execute('SELECT * FROM companies WHERE id=?', (company_id,)).fetchone()

    new = dict(row)
    changes = []
    for f in FIELDS:
        ov = old.get(f) or ''
        nv = new.get(f) or ''
        if str(ov) != str(nv):
            changes.append(f'{FIELD_LABELS.get(f,f)}: "{ov}" → "{nv}"')
    change_str = ' | '.join(changes) if changes else 'No field changes detected'
    bump_version()
    log_action('Company Updated', f'Updated "{new.get("company_name")}" ({new.get("country","UAE")}). Changes: {change_str}')
    return jsonify(new)


@app.route('/api/companies/<int:company_id>', methods=['DELETE'])
@require_admin
def delete_company(company_id):
    with get_db() as conn:
        row = conn.execute('SELECT * FROM companies WHERE id=?', (company_id,)).fetchone()
        conn.execute('DELETE FROM companies WHERE id=?', (company_id,))
        conn.commit()
    bump_version()
    if row:
        r = dict(row)
        parts = [f'Country: {r.get("country") or "UAE"}']
        if r.get('cat'): parts.append(f'Category: {r["cat"]}')
        if r.get('activity'): parts.append(f'Activity: {r["activity"]}')
        if r.get('dnb_rating'): parts.append(f'D&B Rating: {r["dnb_rating"]}')
        if r.get('dnb_rep_validity'): parts.append(f'D&B Validity: {r["dnb_rep_validity"]}')
        if r.get('audit_rep'): parts.append(f'Last Audit: {r["audit_rep"]}')
        log_action('Company Deleted', f'Permanently deleted "{r["company_name"]}". {" | ".join(parts)}')
    else:
        log_action('Company Deleted', f'Deleted company ID {company_id}')
    return '', 204


# ── Excel export ────────────────────────────────────────────────────────────

EXCEL_HEADERS = [
    'SR.No.', 'Cat', 'Company Name', 'Activity', 'Country',
    'D&B Rep Validity', 'D&B Rating', 'D&B Processed By',
    'Last Audit', 'Audit Processed By', 'Auditor Name',
    'Mgtm A/c Q1', 'Mgtm A/c Q2', 'Mgtm A/c Q3', 'Mgtm A/c Q4',
    'AECB (Dir)', 'AECB (Com)', 'Cibil (Dir)',
    'Remarks',
]

EXCEL_FIELDS = [
    None,  # SR.No. — computed
    'cat', 'company_name', 'activity', 'country',
    'dnb_rep_validity', 'dnb_rating', 'dnb_processed_by',
    'audit_rep', 'audit_processed_by', 'audit_prepared_by',
    'mgtm_ac_q1', 'mgtm_ac_q2', 'mgtm_ac_q3', 'mgtm_ac_q4',
    'aecb_dir', 'aecb_com', 'cibil_dir',
    'remarks',
]

# Column widths (one per column, A–S)
COL_WIDTHS = [8, 6, 35, 25, 10, 16, 13, 18, 13, 18, 18, 13, 13, 13, 13, 13, 13, 13, 45]

LAST_COL = get_column_letter(len(EXCEL_HEADERS))  # 'S'


def _build_workbook(companies):
    """Build and return an openpyxl Workbook for the given list of companies."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'D & B'

    title_font  = Font(name='Calibri', bold=True, size=12)
    header_font = Font(name='Calibri', bold=True, size=10)
    sub_font    = Font(name='Calibri', bold=True, size=10)
    data_font   = Font(name='Calibri', size=10)

    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left   = Alignment(horizontal='left',   vertical='center', wrap_text=True)

    thin = Side(style='thin')
    bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)

    hdr_fill = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
    sub_fill = PatternFill(start_color='FFFFCC', end_color='FFFFCC', fill_type='solid')

    ws.merge_cells(f'B1:{LAST_COL}1')
    ws['B1'].value = 'Corporate Credit Information (MIS)'
    ws['B1'].font  = title_font
    ws['B1'].alignment = center

    ws.merge_cells(f'B2:{LAST_COL}2')
    ws['B2'].value = f'({datetime.now().strftime("%d-%m-%Y")})'
    ws['B2'].alignment = center

    for col_i, h in enumerate(EXCEL_HEADERS, 1):
        c = ws.cell(row=3, column=col_i, value=h)
        c.font      = header_font
        c.fill      = hdr_fill
        c.alignment = center
        c.border    = bdr

    uae    = [c for c in companies if (c.get('country') or '').strip().upper() == 'UAE']
    others: dict = {}
    for c in companies:
        ctry = (c.get('country') or '').strip()
        if ctry.upper() != 'UAE' and ctry:
            others.setdefault(ctry, []).append(c)

    row_num = 4

    def write_group(group, label=None):
        nonlocal row_num
        sr = 1
        if label:
            ws.merge_cells(f'A{row_num}:{LAST_COL}{row_num}')
            lc = ws.cell(row=row_num, column=1, value=label)
            lc.font      = sub_font
            lc.fill      = sub_fill
            lc.alignment = center
            row_num += 1

        for comp in group:
            for col_i, field in enumerate(EXCEL_FIELDS, 1):
                val = sr if field is None else (comp.get(field) or '')
                cell = ws.cell(row=row_num, column=col_i, value=val)
                cell.font      = data_font
                cell.alignment = center if col_i <= 2 else left
                cell.border    = bdr
            row_num += 1
            sr += 1

    write_group(uae)
    for country_name, group in others.items():
        write_group(group, label=f'— {country_name} —')

    for i, w in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 22
    ws.row_dimensions[3].height = 32

    return wb


@app.route('/export')
@require_auth
def export_excel():
    with get_db() as conn:
        rows = conn.execute(
            'SELECT * FROM companies ORDER BY country, cat, sort_order, id'
        ).fetchall()
    companies = [dict(r) for r in rows]

    buf = io.BytesIO()
    _build_workbook(companies).save(buf)
    buf.seek(0)

    filename = f'MIS_{datetime.now().strftime("%d-%m-%Y")}.xlsx'
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/export/<int:company_id>')
@require_auth
def export_single(company_id):
    with get_db() as conn:
        row = conn.execute('SELECT * FROM companies WHERE id=?', (company_id,)).fetchone()
    if not row:
        return 'Not found', 404

    company = dict(row)
    buf = io.BytesIO()
    _build_workbook([company]).save(buf)
    buf.seek(0)

    safe = ''.join(c for c in (company.get('company_name') or 'company') if c.isalnum() or c in ' _-')
    filename = f'MIS_{safe}_{datetime.now().strftime("%d-%m-%Y")}.xlsx'
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


if __name__ == '__main__':
    init_db()
    port  = int(os.environ.get('PORT', 5001))
    debug = os.environ.get('DEBUG', 'False').lower() == 'true'
    app.run(debug=debug, port=port)
